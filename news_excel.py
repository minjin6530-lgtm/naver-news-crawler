"""
Excel 서식 생성 (3단계)
입력: news_filtered.csv
결과: news_summary_YYYYMMDD.xlsx

필요 패키지: pip install openpyxl
"""

import csv
import os
from datetime import datetime

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    raise ImportError("openpyxl을 설치하세요: pip install openpyxl")

INPUT  = "news_filtered.csv"
TODAY  = datetime.today().strftime("%Y%m%d")
OUTPUT = f"news_summary_{TODAY}.xlsx"

# ── 서식 상수 ──────────────────────────────────────────
HEADER_FONT  = Font(name="Arial", bold=True, color="FFFFFF", size=10)
HEADER_FILL  = PatternFill("solid", fgColor="305496")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

BODY_FONT    = Font(name="Arial", size=10)
BODY_ALIGN   = Alignment(wrap_text=True, vertical="top")

LINK_FONT    = Font(name="Arial", size=10, color="0563C1", underline="single")

THIN_BORDER  = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)

# 컬럼 정의: (헤더명, 필드키, 열 너비)
COLUMNS = [
    ("검색 키워드",  "search_keyword", 16),
    ("발행일",       "pubDate",         12),
    ("제목",         "title_korean",    42),
    ("요약",         "summary_korean",  50),
    ("중요도",       "importance",      10),
    ("언론사",       "source",          14),
    ("링크",         "link",            38),
]

# importance별 배경색 (연한 하이라이트)
IMPORTANCE_FILL = {
    "★★★": PatternFill("solid", fgColor="FFF2CC"),  # 연노랑
    "★★":  PatternFill("solid", fgColor="EDEDED"),  # 연회색
    "★":   None,
}


def read_csv(path: str) -> list:
    with open(path, newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def main():
    rows = read_csv(INPUT)
    print(f"입력: {len(rows)}행")

    wb = Workbook()
    ws = wb.active
    ws.title = f"뉴스_{TODAY}"

    # ── 헤더 행 ──────────────────────────────────────────
    ws.row_dimensions[1].height = 28
    for col_idx, (header, _, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font   = HEADER_FONT
        cell.fill   = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # ── 본문 행 ──────────────────────────────────────────
    for row_idx, data in enumerate(rows, start=2):
        ws.row_dimensions[row_idx].height = 60

        importance = data.get("importance", "★")
        row_fill = IMPORTANCE_FILL.get(importance)

        for col_idx, (_, field, _) in enumerate(COLUMNS, start=1):
            value = data.get(field, "")
            cell  = ws.cell(row=row_idx, column=col_idx)

            # 링크 컬럼(마지막)은 하이퍼링크 처리
            if field == "link" and value:
                cell.value     = value
                cell.hyperlink = value
                cell.font      = LINK_FONT
                cell.alignment = BODY_ALIGN
            else:
                cell.value     = value
                cell.font      = BODY_FONT
                cell.alignment = BODY_ALIGN

            cell.border = THIN_BORDER
            if row_fill:
                cell.fill = row_fill

    # ── 틀 고정 + 자동 필터 ──────────────────────────────
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(OUTPUT)

    # ── 요약 보고 ─────────────────────────────────────────
    print(f"\n{'='*50}")
    print(f"저장 완료 → {OUTPUT}")
    print(f"총 행 수: {len(rows)}행")

    # 빈 셀 체크 (주요 컬럼)
    for field in ["title_korean", "summary_korean", "importance", "source"]:
        empty = sum(1 for r in rows if not r.get(field, "").strip())
        if empty:
            print(f"  ⚠ '{field}' 빈 셀: {empty}개")

    # importance 분포
    dist = {}
    for r in rows:
        k = r.get("importance", "?")
        dist[k] = dist.get(k, 0) + 1
    print("\nimportance 분포:")
    for k in sorted(dist, reverse=True):
        bar = "█" * dist[k]
        print(f"  {k:4s}  {dist[k]:3d}건  {bar}")
    print(f"{'='*50}")


if __name__ == "__main__":
    main()
